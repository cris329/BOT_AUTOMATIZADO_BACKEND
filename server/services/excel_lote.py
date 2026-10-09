"""
Leo Excel/CSV de documentos o de inscripciones (plantillas simples).

@author Cristian Deysdayr Jiménez
@created 2026-10-09
"""

from __future__ import annotations

import csv
import io
from typing import BinaryIO

from openpyxl import Workbook, load_workbook

from app.types import DocumentoLote


def parse_documentos(file: BinaryIO, filename: str) -> list[DocumentoLote]:
    """Leo número y tipo opcional desde xlsx o csv."""
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return _docs_csv(file)
    return _docs_xlsx(file)


def parse_inscripciones(file: BinaryIO, filename: str) -> list[dict[str, str]]:
    """Leo documento, programa y tipo opcional."""
    name = (filename or "").lower()
    rows = _rows_csv(file) if name.endswith(".csv") else _rows_xlsx(file)
    out: list[dict[str, str]] = []
    for row in rows:
        doc = (row[0] if len(row) > 0 else "").strip()
        prog = (row[1] if len(row) > 1 else "").strip()
        tipo = (row[2] if len(row) > 2 else "").strip()
        if not doc or not prog or _es_encabezado(doc):
            continue
        out.append({"numero_documento": doc, "programa": prog, "tipo_documento": tipo})
    return out


def plantilla_registro_xlsx() -> bytes:
    """Plantilla Fase 1: documento + tipo opcional."""
    return _wb_bytes([["numero_documento", "tipo_documento"], ["1234567890", "CC"]])


def plantilla_inscripciones_xlsx() -> bytes:
    """Plantilla Fase 2: documento + programa + tipo."""
    return _wb_bytes(
        [
            ["numero_documento", "programa", "tipo_documento"],
            ["1234567890", "ANALISIS Y DESARROLLO DE SOFTWARE", "CC"],
        ]
    )


def _docs_xlsx(file: BinaryIO) -> list[DocumentoLote]:
    return [
        DocumentoLote(numero_documento=r[0], tipo_documento=r[1] if len(r) > 1 else "")
        for r in _rows_xlsx(file)
        if r and r[0] and not _es_encabezado(r[0])
    ]


def _docs_csv(file: BinaryIO) -> list[DocumentoLote]:
    return [
        DocumentoLote(numero_documento=r[0], tipo_documento=r[1] if len(r) > 1 else "")
        for r in _rows_csv(file)
        if r and r[0] and not _es_encabezado(r[0])
    ]


def _rows_xlsx(file: BinaryIO) -> list[list[str]]:
    wb = load_workbook(file, read_only=True, data_only=True)
    sheet = wb.active
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any(cells):
            rows.append(cells)
    return rows


def _rows_csv(file: BinaryIO) -> list[list[str]]:
    text = io.TextIOWrapper(file, encoding="utf-8-sig", newline="")
    return [[(c or "").strip() for c in row] for row in csv.reader(text)]


def _es_encabezado(valor: str) -> bool:
    v = valor.strip().lower()
    return v in {"numero_documento", "documento", "número", "numero"}


def _wb_bytes(rows: list[list[str]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
